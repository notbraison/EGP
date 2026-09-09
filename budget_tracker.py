## budget tracker
from copy import copy
import json
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Border, Font, PatternFill
import urllib.request
from scrape_entities import normalize_key, safe_save_workbook
from sync_gsuite import BUDGET_SHEET_URL

EXCEL_FILE = "Procuring_Entities.xlsx"
BUDGET_SHEET = "Budget Totals"
ENTITIES_SHEET = "latest_entities"
CURRENCY_FORMAT = '"KES" #,##0.00'


def pull_status_from_gsheet(file_path=EXCEL_FILE, sheet_url=None):
    """
    Google Sheet is the source of truth for Status + Prequal (columns D/E)
    only — your collaborator edits there. Total Budget stays locally
    computed. Run this BEFORE budget_tracker/reconcile so local processing
    never clobbers his edits; push (sync_gsuite) runs AFTER, so his status
    round-trips back up unchanged alongside the freshly computed totals.
    """
    url = sheet_url or BUDGET_SHEET_URL
    try:
        with urllib.request.urlopen(url, timeout=30) as resp:
            payload = json.loads(resp.read().decode())
    except Exception as e:
        print(f"[ERROR] Could not fetch Google Sheet: {e}")
        return

    rows = payload.get("data", [])[1:]  # skip header row
    remote_map = {}
    for row in rows:
        if len(row) < 5 or not row[1]:
            continue
        remote_map[normalize_key(row[1])] = (row[3], row[4])

    try:
        wb = load_workbook(file_path)
    except PermissionError:
        print(f"[PERMISSION ERROR] '{file_path}' is open in Excel.")
        return

    ws = wb[BUDGET_SHEET]
    updated = 0
    for r in range(2, ws.max_row + 1):
        entity = ws.cell(row=r, column=2).value
        if not entity:
            continue
        key = normalize_key(entity)
        if key in remote_map:
            r_status, r_prequal = remote_map[key]
            if ws.cell(row=r, column=4).value != r_status or ws.cell(row=r, column=5).value != r_prequal:
                ws.cell(row=r, column=4, value=r_status)
                ws.cell(row=r, column=5, value=r_prequal)
                updated += 1

    safe_save_workbook(wb, file_path)
    print(f"Pulled Status/Prequal from Sheet — {updated} rows updated.")

def update_budget_totals(file_path=EXCEL_FILE):
    pull_status_from_gsheet(file_path)
    wb = openpyxl.load_workbook(file_path)

    if BUDGET_SHEET not in wb.sheetnames or ENTITIES_SHEET not in wb.sheetnames:
        raise ValueError(
            f"Workbook must contain both '{BUDGET_SHEET}' and"
            f" '{ENTITIES_SHEET}'"
        )

    s1 = wb[BUDGET_SHEET]
    s2 = wb[ENTITIES_SHEET]

    # Column order: 1 Sr.No | 2 Entity | 3 Total Budget | 4 Status | 5 Prequalification done
    existing_map = {}
    for r in range(2, s1.max_row + 1):
        entity = s1.cell(row=r, column=2).value
        budget = s1.cell(row=r, column=3).value
        status = s1.cell(row=r, column=4).value
        prequal = s1.cell(row=r, column=5).value

        if entity:
            clean_entity = str(entity).replace("\xa0", " ").strip()
            key = "".join(e for e in clean_entity.upper() if e.isalnum())
            existing_map[key] = {
                "budget": budget,
                "status": status,
                "prequal": prequal,
            }

    new_entities = []
    for r in range(2, s2.max_row + 1):
        sr_no = s2.cell(row=r, column=1).value
        entity = s2.cell(row=r, column=3).value

        if entity:
            clean_entity = str(entity).replace("\xa0", " ").strip()
            new_entities.append((sr_no, clean_entity))

    # Template cells taken from the first real data row (row 2)
    template_row = [s1.cell(row=2, column=c) for c in range(1, 6)]

    for i, (sr_no, entity_name) in enumerate(new_entities):
        row_idx = i + 2
        key = "".join(e for e in entity_name.upper() if e.isalnum())
        matched = existing_map.get(key, {"budget": 0.0, "status": "Pending", "prequal": ""})

        status_val = (
            str(matched["status"]).strip().lower() if matched["status"] else ""
        )

        s1.cell(row=row_idx, column=1).value = sr_no or (i + 1)
        s1.cell(row=row_idx, column=2).value = entity_name
        s1.cell(row=row_idx, column=3).value = matched.get("budget", 0.0)
        s1.cell(row=row_idx, column=4).value = matched.get("status", "Pending")
        s1.cell(row=row_idx, column=5).value = matched.get("prequal", "")

        for c in range(1, 6):
            cell = s1.cell(row=row_idx, column=c)
            tmpl = template_row[c - 1]

            cell.border = copy(tmpl.border)
            cell.alignment = copy(tmpl.alignment)
            cell.font = copy(tmpl.font)
            if c != 3:
                cell.number_format = copy(tmpl.number_format)

        # Total Budget always formatted as KES currency, 2dp
        s1.cell(row=row_idx, column=3).number_format = CURRENCY_FORMAT

        # Status styling applied to Column 4
        status_cell = s1.cell(row=row_idx, column=4)
        if status_val == "done":
            status_cell.fill = PatternFill(
                start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
            )
            status_cell.font = Font(
                name=template_row[3].font.name or "Arial",
                size=template_row[3].font.size or 11,
                bold=True,
                color="375623",
            )
        elif status_val == "partial":
            status_cell.fill = PatternFill(
                start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
            )
            status_cell.font = Font(
                name=template_row[3].font.name or "Arial",
                size=template_row[3].font.size or 11,
                bold=True,
                color="9C6500",
            )
        else:
            status_cell.fill = PatternFill(fill_type=None)
            status_cell.font = Font(
                name=template_row[3].font.name or "Arial",
                size=template_row[3].font.size or 11,
                bold=False,
                color="000000",
            )

    # Clean up orphan trailing rows across all 5 columns
    last_valid_row = len(new_entities) + 1
    if s1.max_row > last_valid_row:
        for r in range(last_valid_row + 1, s1.max_row + 1):
            for c in range(1, 6):
                cell = s1.cell(row=r, column=c)
                cell.value = None
                cell.border = Border()
                cell.fill = PatternFill(fill_type=None)

    wb.save(file_path)
    print(f"Updated '{BUDGET_SHEET}' with {len(new_entities)} entities.")


if __name__ == "__main__":
    update_budget_totals()