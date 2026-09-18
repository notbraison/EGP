## sync_gsuite.py
import json
import urllib.request
import openpyxl

EXCEL_FILE = "Procuring_Entities.xlsx"

# Web App URLs for both Google Sheets
BUDGET_SHEET_URL = "https://script.google.com/macros/s/AKfycbxK0LfVlhReWLp6dwmYscZa-Y_gm8zQRqbeoFq7xXOP3QnXzr_NeOpa2FCnrXxDjVGP/exec"
VISITS_SHEET_URL = "https://script.google.com/macros/s/AKfycbzRjui20V9982w3Rdr8LiLJJJ3xOVACCfhZR-T50dCzAKbFGlKi810hostkA5tV91CceA/exec"


def sync_sheet_to_url(sheet_name, web_app_url, num_cols, remote_sheet_name=None):
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    if sheet_name not in wb.sheetnames:
        print(f"[SKIP] Sheet '{sheet_name}' not found in Excel.")
        return

    s = wb[sheet_name]
    data = []

    for row in s.iter_rows(values_only=True):
        if any(cell is not None for cell in row):
            clean_row = [
                cell if cell is not None else "" for cell in row[:num_cols]
            ]
            data.append(clean_row)

    if not data:
        print(f"No data found in '{sheet_name}'.")
        return

    payload = json.dumps({
        "data": data,
        "sheetName": remote_sheet_name or sheet_name,
    }).encode("utf-8")
    req = urllib.request.Request(
        web_app_url,
        data=payload,
        headers={"Content-Type": "application/json"},
        method="POST",
    )

    print(f"Syncing '{sheet_name}' to Google Sheets...")
    with urllib.request.urlopen(req) as response:
        result = json.loads(response.read().decode())
        if result.get("status") == "success":
            print(f"Successfully synced '{sheet_name}' ({len(data)} rows).")
        else:
            print(f"Failed to sync '{sheet_name}'.")


def sync_all_google_sheets():
    sync_sheet_to_url("Budget Totals", BUDGET_SHEET_URL, num_cols=5)
    sync_sheet_to_url("addresses", VISITS_SHEET_URL, num_cols=9)
    sync_sheet_to_url("blacklisted", BUDGET_SHEET_URL, num_cols=2, remote_sheet_name="blacklisted")


if __name__ == "__main__":
    sync_all_google_sheets()