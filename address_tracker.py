## address_tracker.py
from copy import copy
import re
import time
import random
import threading
import urllib.parse
from concurrent.futures import ThreadPoolExecutor, as_completed
import openpyxl
from openpyxl.styles import Border, PatternFill
from playwright.sync_api import sync_playwright
from scrape_entities import is_blacklisted

EXCEL_FILE = "Procuring_Entities.xlsx"
ENTITIES_SHEET = "latest_entities"
ADDRESSES_SHEET = "addresses"

MAX_WORKERS = 4          # concurrent Bing searches
SAVE_EVERY = 10          # write to disk every N completed searches, not every 1
GOTO_TIMEOUT_MS = 15000
SELECTOR_TIMEOUT_MS = 4000
MAX_RETRIES = 2
RETRY_SLEEP = 3
MIN_DELAY = 1.2
MAX_DELAY = 3.0

HEADERS = [
    "Sr. No.", "Procuring Entity", "Physical Address (Building/Street)",
    "Town/Area", "County", "Phone", "Email", "ASSIGNED", "VISIT"
]

KENYA_COUNTIES = [
    "Nairobi", "Uasin Gishu", "Nyandarua", "Kiambu", "Nandi", "Nyeri",
    "Kirinyaga", "Busia", "Kisumu", "Bungoma", "Machakos", "Migori",
    "Bomet", "West Pokot", "Baringo", "Kakamega", "Homa Bay", "Mombasa",
    "Kajiado", "Nakuru", "Murang'a", "Trans Nzoia", "Kilifi", "Kwale",
    "Taita Taveta", "Garissa", "Wajir", "Mandera", "Marsabit", "Isiolo",
    "Meru", "Tharaka-Nithi", "Embu", "Kitui", "Makueni", "Turkana",
    "Samburu", "Elgeyo Marakwet", "Kericho", "Laikipia", "Narok", "Vihiga",
    "Siaya", "Kisii", "Nyamira", "Tana River", "Lamu",
]


def clean_addresses_sheet(file_path=EXCEL_FILE):
    """Optional standalone utility — strips any blacklisted rows already
    present in 'addresses'. Not required in the normal flow anymore, since
    update_address_tracker() now filters blacklisted entities before they
    ever get written, but kept here in case of stale data from an older run.
    """
    wb = openpyxl.load_workbook(file_path)
    if "addresses" not in wb.sheetnames:
        return

    ws = wb["addresses"]
    rows_to_keep = []
    headers = [cell.value for cell in ws[1]]
    rows_to_keep.append(headers)

    for row in ws.iter_rows(min_row=2, values_only=True):
        entity_name = row[1]
        if entity_name and not is_blacklisted(entity_name):
            rows_to_keep.append(row)

    ws.delete_rows(1, ws.max_row)
    for r in rows_to_keep:
        ws.append(r)

    wb.save(file_path)


def normalize_key(text):
    if not text:
        return ""
    clean = str(text).replace("\xa0", " ").strip().upper()
    return "".join(c for c in clean if c.isalnum())


def extract_contact_info(search_text):
    info = {"address": "", "town": "", "county": "", "phone": "", "email": ""}
    if not search_text:
        return info

    email_match = re.search(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}", search_text)
    if email_match:
        info["email"] = email_match.group(0).lower()

    phone_match = re.search(r"(?:\+?254|0)(?:\s?\d{2,3}){1,3}\s?\d{3,4}\b", search_text)
    if phone_match:
        raw_phone = phone_match.group(0).strip()
        info["phone"] = f"'{raw_phone}" if not raw_phone.startswith("'") else raw_phone

    for county in KENYA_COUNTIES:
        if re.search(r"\b" + re.escape(county) + r"\b", search_text, re.IGNORECASE):
            info["county"] = county
            break

    road_match = re.search(
        r"([A-Za-z0-9\s,-]+(?:Road|Street|Avenue|Way|Building|House|Towers|Plaza|Campus))",
        search_text, re.IGNORECASE,
    )
    if road_match:
        info["address"] = road_match.group(0).strip()

    return info


def perform_single_search(entity, page):
    """Executes a single web search using Bing."""
    clean_query = re.sub(r'^\d+\s*-?\s*', '', entity)
    query = f'"{clean_query}" physical address contact phone email county Kenya'
    search_url = f"https://www.bing.com/search?q={urllib.parse.quote(query)}"

    for attempt in range(MAX_RETRIES):
        try:
            page.goto(search_url, wait_until="domcontentloaded", timeout=GOTO_TIMEOUT_MS)
            try:
                page.wait_for_selector("#b_results", timeout=SELECTOR_TIMEOUT_MS)
            except Exception:
                pass
            combined_text = page.evaluate("document.body.innerText")
            return extract_contact_info(combined_text)
        except Exception as e:
            print(f"    [Attempt {attempt + 1}/{MAX_RETRIES}] Failed search for {entity}: {e}")
            if attempt < MAX_RETRIES - 1:
                time.sleep(RETRY_SLEEP)

    return {"address": "", "town": "", "county": "", "phone": "", "email": ""}


def _worker_process_chunk(chunk, worker_id):
    """Runs entirely on its own thread with its own Playwright/browser instance.
    Playwright's sync API is not safe to share across threads — this is why
    each worker must own its full browser lifecycle rather than borrowing a
    page created elsewhere.
    """
    results = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        )
        page = context.new_page()
        page.route("**/*.{png,jpg,jpeg,svg,gif,woff,css}", lambda route: route.abort())

        for row_idx, entity_name in chunk:
            result = perform_single_search(entity_name, page)
            results.append((row_idx, entity_name, result))
            time.sleep(random.uniform(MIN_DELAY, MAX_DELAY))

        browser.close()
    return results


def update_address_tracker(file_path=EXCEL_FILE):
    wb = openpyxl.load_workbook(file_path)

    if ENTITIES_SHEET not in wb.sheetnames:
        raise ValueError(f"'{ENTITIES_SHEET}' must exist with scraped entity data.")
    s2 = wb[ENTITIES_SHEET]

    if ADDRESSES_SHEET not in wb.sheetnames:
        s3 = wb.create_sheet(ADDRESSES_SHEET)
        s3.append(HEADERS)
    else:
        s3 = wb[ADDRESSES_SHEET]

    # 1. Index existing address records
    existing_data = {}
    if s3.max_row >= 2:
        for r in range(2, s3.max_row + 1):
            entity = s3.cell(row=r, column=2).value
            if entity:
                key = normalize_key(entity)
                existing_data[key] = {
                    "address": s3.cell(row=r, column=3).value or "",
                    "town": s3.cell(row=r, column=4).value or "",
                    "county": s3.cell(row=r, column=5).value or "",
                    "phone": s3.cell(row=r, column=6).value or "",
                    "email": s3.cell(row=r, column=7).value or "",
                    "assigned": s3.cell(row=r, column=8).value or "",
                    "visit": s3.cell(row=r, column=9).value or "",
                }

    # 2. Build entity list, skipping blacklisted entities entirely — no point
    # burning a Bing search on something we're going to exclude anyway.
    scraped_entities = []
    skipped_blacklisted = 0
    for r in range(2, s2.max_row + 1):
        entity = s2.cell(row=r, column=3).value
        if entity:
            clean = str(entity).replace("\xa0", " ").strip()
            if is_blacklisted(clean):
                skipped_blacklisted += 1
                continue
            scraped_entities.append(clean)

    if skipped_blacklisted:
        print(f"Skipped {skipped_blacklisted} blacklisted entities (no search performed).")

    template_cells = [s3.cell(row=2, column=c) for c in range(1, 10)] if s3.max_row >= 2 else None

    # 3. Split into "already complete" (write immediately) vs "needs search"
    row_plan = {}       # row_idx -> row_values (already known)
    to_search = []      # (row_idx, entity_name)

    for i, entity_name in enumerate(scraped_entities, start=1):
        row_idx = i + 1
        key = normalize_key(entity_name)
        record = existing_data.get(key, {})

        address = record.get("address", "")
        town = record.get("town", "")
        county = record.get("county", "")
        phone = record.get("phone", "")
        email = record.get("email", "")

        needs_search = not all([address, phone, email, county])

        if needs_search:
            to_search.append((row_idx, entity_name))
        else:
            if phone and not str(phone).startswith("'"):
                phone = f"'{phone}"
            row_plan[row_idx] = [
                i, entity_name, address, town, county, phone, email,
                record.get("assigned", ""), record.get("visit", ""),
            ]

    write_lock = threading.Lock()
    save_lock = threading.Lock()
    completed_count = 0

    def write_row(row_idx, i, entity_name, address, town, county, phone, email, assigned, visit):
        if phone and not str(phone).startswith("'"):
            phone = f"'{phone}"
        row_values = [i, entity_name, address, town, county, phone, email, assigned, visit]
        with write_lock:
            for col_idx, val in enumerate(row_values, start=1):
                cell = s3.cell(row=row_idx, column=col_idx)
                cell.value = val
                if template_cells:
                    tmpl = template_cells[col_idx - 1]
                    cell.border = copy(tmpl.border)
                    cell.alignment = copy(tmpl.alignment)
                    cell.font = copy(tmpl.font)

    # 4. Write the "already complete" rows immediately (no search needed)
    for row_idx, values in row_plan.items():
        with write_lock:
            for col_idx, val in enumerate(values, start=1):
                cell = s3.cell(row=row_idx, column=col_idx)
                cell.value = val
                if template_cells:
                    tmpl = template_cells[col_idx - 1]
                    cell.border = copy(tmpl.border)
                    cell.alignment = copy(tmpl.alignment)
                    cell.font = copy(tmpl.font)

        # 5. Run searches concurrently — each worker owns its own browser instance
    if to_search:
        print(f"Searching {len(to_search)} entities using {MAX_WORKERS} concurrent workers...")

        # Split into MAX_WORKERS roughly-even chunks
        chunks = [[] for _ in range(MAX_WORKERS)]
        for idx, item in enumerate(to_search):
            chunks[idx % MAX_WORKERS].append(item)
        chunks = [c for c in chunks if c]  # drop empty chunks if fewer items than workers

        with ThreadPoolExecutor(max_workers=len(chunks)) as executor:
            futures = {
                executor.submit(_worker_process_chunk, chunk, i): i
                for i, chunk in enumerate(chunks)
            }

            for future in as_completed(futures):
                worker_id = futures[future]
                try:
                    chunk_results = future.result()
                except Exception as e:
                    print(f"  [WORKER {worker_id} ERROR] {e}")
                    continue

                for row_idx, entity_name, searched_data in chunk_results:
                    key = normalize_key(entity_name)
                    record = existing_data.get(key, {})

                    address = record.get("address", "") or searched_data.get("address", "")
                    town = record.get("town", "") or searched_data.get("town", "")
                    county = record.get("county", "") or searched_data.get("county", "")
                    phone = record.get("phone", "") or searched_data.get("phone", "")
                    email = record.get("email", "") or searched_data.get("email", "")
                    i = scraped_entities.index(entity_name) + 1

                    write_row(
                        row_idx, i, entity_name, address, town, county, phone, email,
                        record.get("assigned", ""), record.get("visit", ""),
                    )

                    completed_count += 1
                    print(f"  [{completed_count}/{len(to_search)}] Done: {entity_name}")

                    if completed_count % SAVE_EVERY == 0:
                        with save_lock:
                            try:
                                wb.save(file_path)
                            except Exception as e:
                                print(f"    [!] Could not save file (is it open in Excel?): {e}")

    # 6. Clear dead rows at the bottom
    last_valid_row = len(scraped_entities) + 1
    if s3.max_row > last_valid_row:
        for r in range(last_valid_row + 1, s3.max_row + 1):
            for c in range(1, 10):
                cell = s3.cell(row=r, column=c)
                cell.value = None
                cell.border = Border()
                cell.fill = PatternFill(fill_type=None)

    wb.save(file_path)
    print(f"\nUpdated '{ADDRESSES_SHEET}' with {len(scraped_entities)} entities "
          f"({len(to_search)} searched, {len(row_plan)} already complete).")


if __name__ == "__main__":
    update_address_tracker()