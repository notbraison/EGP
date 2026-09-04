import openpyxl
from copy import copy
from openpyxl.styles import Font, PatternFill, Border

def update_workflow_with_formatting(file_path):
    wb = openpyxl.load_workbook(file_path)
    
    if 'Sheet1' not in wb.sheetnames or 'Sheet2' not in wb.sheetnames:
        raise ValueError("Workbook must contain both 'Sheet1' and 'Sheet2'")

    s1 = wb['Sheet1']
    s2 = wb['Sheet2']

    # 1. Index existing budgets and statuses from Sheet1 (Data starts at Row 3)
    existing_map = {}
    for r in range(3, s1.max_row + 1):
        entity = s1.cell(row=r, column=2).value
        budget = s1.cell(row=r, column=3).value
        status = s1.cell(row=r, column=4).value
        
        if entity:
            clean_entity = str(entity).replace('\xa0', ' ').strip()
            key = "".join(e for e in clean_entity.upper() if e.isalnum())
            existing_map[key] = {
                'budget': budget,
                'status': status
            }

    # 2. Extract updated entity order from Sheet2 (Data starts at Row 2)
    new_entities = []
    for r in range(2, s2.max_row + 1):
        sr_no = s2.cell(row=r, column=1).value
        entity = s2.cell(row=r, column=3).value  # Entity name sits in Column C of Sheet2
        
        if entity:
            clean_entity = str(entity).replace('\xa0', ' ').strip()
            new_entities.append((sr_no, clean_entity))

    # Reference template row (Row 3) for standard cell styles & borders
    template_row = [s1.cell(row=3, column=c) for c in range(1, 5)]

    # 3. Update values and apply formatting across all rows
    for i, (sr_no, entity_name) in enumerate(new_entities):
        row_idx = i + 3
        key = "".join(e for e in entity_name.upper() if e.isalnum())
        matched = existing_map.get(key, {'budget': '', 'status': ''})
        
        status_val = str(matched['status']).strip().lower() if matched['status'] else ''

        # Set values
        s1.cell(row=row_idx, column=1).value = sr_no or (i + 1)
        s1.cell(row=row_idx, column=2).value = entity_name
        s1.cell(row=row_idx, column=3).value = matched['budget']
        s1.cell(row=row_idx, column=4).value = matched['status']

        # Copy base borders, font family, alignment, and number formats
        for c in range(1, 5):
            cell = s1.cell(row=row_idx, column=c)
            tmpl = template_row[c - 1]
            
            cell.border = copy(tmpl.border)
            cell.alignment = copy(tmpl.alignment)
            cell.font = copy(tmpl.font)
            cell.number_format = copy(tmpl.number_format)

        # Apply specific styling for Status (Column 4)
        status_cell = s1.cell(row=row_idx, column=4)
        if status_val == 'done':
            status_cell.fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
            status_cell.font = Font(
                name=template_row[3].font.name or 'Arial', 
                size=template_row[3].font.size or 11, 
                bold=True, 
                color='375623'
            )
        else:
            status_cell.fill = PatternFill(fill_type=None)
            status_cell.font = Font(
                name=template_row[3].font.name or 'Arial', 
                size=template_row[3].font.size or 11, 
                bold=False, 
                color='000000'
            )

    # 4. Clean up trailing rows beyond updated data length
    last_valid_row = len(new_entities) + 2
    if s1.max_row > last_valid_row:
        for r in range(last_valid_row + 1, s1.max_row + 1):
            for c in range(1, 5):
                cell = s1.cell(row=r, column=c)
                cell.value = None
                cell.border = Border()
                cell.fill = PatternFill(fill_type=None)

    wb.save(file_path)
    print(f"Updated {len(new_entities)} entities in place with complete formatting.")

if __name__ == '__main__':
    update_workflow_with_formatting('Procuring_Entities_Workflow.xlsx')